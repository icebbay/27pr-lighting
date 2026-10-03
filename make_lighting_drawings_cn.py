"""GF lighting construction drawings in Chinese practice (GB/T 50786-2012 / 09DX001 style).

    3D_gen_bench\\_tools_venv\\Scripts\\python.exe products\\make_lighting_drawings_cn.py

Output: Downloads\\27PR_一层照明施工图_RevB_国标.pptx
  电施-01 设计说明 · 图例 · 灯具表
  电施-02 一层照明平面图（全层：配电箱 AL1、回路 WL1~3、导线、开关）
  电施-03~05 分区放大图
  电施-06 AL1 照明配电箱系统图
  电施-07 开关面板表 + 回路表
Reuses the circuit data / geometry from make_lighting_drawings.py (Rev B, common-sense resolutions included).
Chinese plans draw the actual wiring: home run from the board (WL#), power line linking the lamps of a WL,
switch drops, conductor counts as slashes; the UK set uses circuit letters instead.
"""
import math, copy, os
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

_path = __file__.replace("make_lighting_drawings_cn.py", "make_lighting_drawings.py")
_src = open(_path, encoding="utf-8").read()
_src = _src[:_src.index("# ---------------- key plan ----------------")]
G = {"__file__": _path}
exec(compile(_src, _path, "exec"), G)
globals().update({k: v for k, v in G.items() if not k.startswith("__")})
DST = rf"C:\Users\jcjia\Downloads\27PR_{FL_CN}照明施工图_RevB_国标.pptx"
DST = os.environ.get("DST_OVERRIDE", DST)
SB = 10 if FLOOR == "FF" else 0
NZ = len(ZONES)
def SN(i): return f"电施-{SB + i:02d}"   # GF 电施-01..07, FF 电施-11..16
N0 = len(p.slides)

# ---------------- circuits (配电回路) ----------------
WL = WL_GF_DEF = {"WL1": dict(name="一层前区照明（门厅 / 起居室 / G07 区 / 楼梯）", circ=["L1", "L2", "L4", "L5", "L6", "L7", "L8", "L9", "L14"], br="C16 1P+N RCBO 30mA", n=21),
      "WL2": dict(name="一层后区照明（G15 区 / 卫生间 / 厨房 / 洗衣房）+ 排气扇", circ=["L10", "L11", "L12", "L13", "L15", "L18", "L19", "L20"], br="C16 1P+N RCBO 30mA", n=14),
      "WL3": dict(name="户外照明（门廊 / 后花园壁灯 / 围栏地灯）", circ=["L3", "L16", "L17", "L21"], br="C10 1P+N RCBO 30mA", n=8)}
WL_FF = {"WL4": dict(name="二层西区照明（主卧 / 盥洗室 / 西卫 / 衣帽间 / 卧室2）+ 排气扇", circ=["L1", "L2", "L6", "L7", "L8", "L9", "L10"], br="C16 1P+N RCBO 30mA", n=9),
         "WL5": dict(name="二层东区照明（走廊 / 北卫 / 卧室3 / 楼梯 / 书房 / 后卫）+ 排气扇", circ=["L3", "L4", "L5", "L11", "L12", "L13"], br="C16 1P+N RCBO 30mA", n=7)}
AL1 = (960 * U, 1012 * U)                  # under the stairs (page-2 coordinates), position to be set on site
BOARD_LABEL = "AL1"
if FLOOR == "FF":
    WL = {k: v for k, v in WL_FF.items()}
    AL1 = (1205 * U, 1062 * U)            # riser at the stair well, from AL1 on the ground floor
    BOARD_LABEL = "↑ 引自一层 AL1"
CN_NOTES = None
if os.environ.get("LIGHTING_JSON"):
    lighting_io.apply_cn(globals())        # WL / AL1 / notes from the web tool
C2W = {c: w for w, d in WL.items() for c in d["circ"]}
LAMP = {n: next(c[0] for c in CIRC if n in c[3]) for c in CIRC for n in c[3]}


def lamp_type(n):
    if n.startswith("主灯位待选"): return "待选主灯"
    if n.startswith("主灯"): return "吊灯/吸顶灯"
    if n.startswith("射灯"): return "筒灯"
    if n.startswith("壁灯"): return "壁灯"
    return "地埋灯"


# ---------------- Chinese symbols ----------------
def cn_light(pen, n, x, y, r):
    t = lamp_type(n)
    if t == "吊灯/吸顶灯": pen.pendant(x, y, r)
    elif t == "待选主灯": pen.pendant(x, y, r, True)
    elif t == "筒灯": pen.down(x, y, r)
    elif t == "壁灯":   # 09DX001: half-filled circle
        pen.shape(MSO_SHAPE.OVAL, x, y, 2.2 * r, 2.2 * r)
        pts = [(x + 1.1 * r * math.cos(a), y + 1.1 * r * math.sin(a)) for a in [math.pi * i / 16 for i in range(17)]]
        pen.poly(pts + [pts[0]], True, fill=INK, lw=0.5)
    else:
        pen.ground(x, y, r)


def cn_switch(pen, x, y, gangs, two_way, r):
    """Filled dot + lever; ticks = gangs; a two-way switch also gets the reverse lever (单极双控)."""
    pen.shape(MSO_SHAPE.OVAL, x, y, 0.9 * r, 0.9 * r, fill=INK, lw=0)
    ang = math.radians(45); L = 2.4 * r
    ex, ey = x + L * math.cos(ang), y - L * math.sin(ang)
    pen.line(x, y, ex, ey, w=1.25)
    for gi in range(max(gangs, 1)):
        t = 1 - gi * 0.18; px, py = x + (ex - x) * t, y + (ey - y) * t
        pen.line(px, py, px + math.sin(ang) * 0.6 * r, py + math.cos(ang) * 0.6 * r, w=1.25)
    if two_way:
        bx, by = x - L * 0.75 * math.cos(ang), y + L * 0.75 * math.sin(ang)
        pen.line(x, y, bx, by, w=1.25); pen.line(bx, by, bx - math.sin(ang) * 0.6 * r, by - math.cos(ang) * 0.6 * r, w=1.25)
    return ex, ey


def board(pen, x, y, r, label):
    pen.shape(MSO_SHAPE.RECTANGLE, x, y, 4.5 * r, 2.2 * r, fill=WHITE, lw=1.0)
    pen.poly([(x - 2.25 * r, y + 1.1 * r), (x + 2.25 * r, y - 1.1 * r), (x + 2.25 * r, y + 1.1 * r)], True, fill=INK, lw=0)
    pen.text(x - 3 * r, y + 1.3 * r, 6 * r, 1.8 * r, label, 7, True, align=PP_ALIGN.CENTER)


def slashes(pen, a, b, n, r):
    """Conductor count on a line: n short slashes at the midpoint (shown only when n != 3, per 09DX001 practice)."""
    if n == 3: return
    (x0, y0), (x1, y1) = a, b; mx, my = (x0 + x1) / 2, (y0 + y1) / 2
    L = math.hypot(x1 - x0, y1 - y0) or 1; ux, uy = (x1 - x0) / L, (y1 - y0) / L
    k = min(n, 5)
    for i in range(k):
        o = (i - (k - 1) / 2) * 0.35 * r; cx, cy = mx + ux * o, my + uy * o
        pen.line(cx - (ux + uy) * 0.45 * r, cy - (uy - ux) * 0.45 * r, cx + (ux + uy) * 0.45 * r, cy + (uy - ux) * 0.45 * r, w=0.75)
    if n > 5:
        pen.text(mx + 0.6 * r, my - 1.6 * r, 3 * r, 1.4 * r, str(n), 6.5, True)


def ortho(pen, a, b, w=0.75, dash=None):
    """Route along the axes (Chinese plans run conduits parallel to walls): a -> corner -> b."""
    (x0, y0), (x1, y1) = a, b
    if abs(x1 - x0) < 1 or abs(y1 - y0) < 1:
        pen.line(x0, y0, x1, y1, w=w, dash=dash); return ((x0 + x1) / 2, (y0 + y1) / 2), (x1 - x0, y1 - y0)
    c = (x1, y0) if abs(x1 - x0) >= abs(y1 - y0) else (x0, y1)
    pen.line(x0, y0, c[0], c[1], w=w, dash=dash); pen.line(c[0], c[1], x1, y1, w=w, dash=dash)
    return ((x0 + c[0]) / 2, (y0 + c[1]) / 2), (c[0] - x0, c[1] - y0)


def ortho_clear(pen, a, b, avoid, clr, w=0.75, dash=None):
    """L-shaped route a->b choosing the corner whose two legs pass farthest from other symbols (avoid)."""
    (x0, y0), (x1, y1) = a, b

    def seg_d(p, q, c):
        (px, py), (qx, qy) = p, q; vx, vy = qx - px, qy - py; L2 = vx * vx + vy * vy or 1
        t = max(0, min(1, ((c[0] - px) * vx + (c[1] - py) * vy) / L2))
        return math.hypot(px + t * vx - c[0], py + t * vy - c[1])

    best = None
    for corner in ((x1, y0), (x0, y1)):
        d = min([seg_d(a, corner, c) for c in avoid] + [seg_d(corner, b, c) for c in avoid] + [9e9])
        if best is None or d > best[0]: best = (d, corner)
    c = best[1]
    pen.line(x0, y0, c[0], c[1], w=w, dash=dash); pen.line(c[0], c[1], x1, y1, w=w, dash=dash)


def mst(pts):
    if len(pts) < 2: return []
    inn, out, edges = [pts[0]], pts[1:], []
    while out:
        a, b = min(((a, b) for a in inn for b in out), key=lambda e: math.hypot(e[0][0] - e[1][0], e[0][1] - e[1][1]))
        edges.append((a, b)); inn.append(b); out.remove(b)
    return edges


def draw_plan(sl, pen, V, zone=None, r=None):
    """Wiring + symbols. zone=None: whole floor; else only that zone's lamps / switches."""
    r = r or 0.11 * M_OLD * V.K
    inz = lambda n: zone is None or LSHEET.get(n) == zone
    swz = lambda pn: zone is None or PSHEET.get(pn) == zone
    # 1) lamps of ONE switched circuit (same letter) are chained together; different letters are never joined
    for cid, desc, keys, lights in CIRC:
        pts = [V.T(LP[n]) for n in lights if inz(n)]
        for a_, b_ in mst(pts):
            ortho(pen, a_, b_, 0.75)
    # 2) feed: AL1 -> switch plates of each WL (thick), chained plate to plate
    for w, d in WL.items():
        plates = [pn for sid, pn, loc, sh in PLATES
                  if any(KEY_CIRC.get(k) in d["circ"] for k in PANEL[pn][1]) and swz(pn)]
        pts = [V.T(MOUNT[pn]) for pn in plates]
        if not pts: continue
        al = V.T(AL1); wi = list(WL).index(w)
        for a_, b_ in mst(pts):
            ortho(pen, a_, b_, 1.25)
        first = min(pts, key=lambda t: math.hypot(t[0] - al[0], t[1] - al[1]))
        if V.inside(al):
            ortho(pen, (al[0] + (wi - 1) * 0.6 * r, al[1]), first, 1.5)
            if zone is not None:
                pen.text(al[0] - 1.5 * r - 1700000, al[1] - (wi + 1.5) * 1.5 * r, 1700000, 150000, f"{w}  BV-3×2.5 PC20 WC/CC", 6.5, True, align=PP_ALIGN.RIGHT)
        else:
            pen.text(first[0] + 1.3 * r, first[1] + 0.8 * r, 1800000, 150000, f"{w} 引自 AL1", 6.5, True)
    # 3) switch drops: one line per key to the nearest lamp of THAT key's circuit (switched live + neutral)
    for sid, pn, loc, sh in PLATES:
        if not swz(pn): continue
        g, keys = PANEL[pn]; m = V.T(MOUNT[pn])
        for k in keys:
            c = KEY_CIRC.get(k)
            if not c: continue
            lamps = [n for n in next(cc for cc in CIRC if cc[0] == c)[3] if inz(n)]
            if lamps:
                t = min((V.T(LP[n]) for n in lamps), key=lambda q: math.hypot(q[0] - m[0], q[1] - m[1]))
                others = [V.T(LP[n]) for n in LP if inz(n) and V.T(LP[n]) != t]
                ortho_clear(pen, m, t, others, r, 0.5, MSO_LINE_DASH_STYLE.SQUARE_DOT)
            else:
                pen.text(m[0] + 0.6 * r, m[1] + 0.8 * r, 1500000, 130000, f"{CL[c]} → 接线见另图", 6.5, color=WIRE)
        for k in keys:
            if k not in RISER: continue
            pen.line(*m, m[0], m[1] - 3 * r, w=0.75)
            if zone is not None: pen.text(m[0] + 0.4 * r, m[1] - 3.6 * r, 2200000, 130000, RISER[k], 6.5)
    # travellers between two-way / intermediate switches
    pairs = {}
    for cid, desc, keys, lights in CIRC:
        pns = [pn for k in keys for pn in KEY2PANEL.get(k, [])]
        for i in range(len(pns) - 1):
            if pns[i] != pns[i + 1]: pairs.setdefault(tuple(sorted((pns[i], pns[i + 1]))), []).append(cid)
    for (a, b), cids in pairs.items():
        if swz(a) and swz(b):
            pa, pb = V.T(MOUNT[a]), V.T(MOUNT[b])
            pen.line(*pa, *pb, w=0.75, dash=MSO_LINE_DASH_STYLE.LONG_DASH)
            slashes(pen, pa, pb, 2 * len(cids) + 1, r)
    # door switches (MK, open door -> light on), wired to the nearest lamp of their circuit
    for ds in DOORSW:
        if zone not in (None, ds["sheet"]): continue
        d = V.T((ds["x"], ds["y"])); pen.box(d[0], d[1], ds["id"], 0.7 * r, 5)
        lamps = next(cc for cc in CIRC if cc[0] == ds["circuit"])[3]
        t = min((V.T(LP[n]) for n in lamps), key=lambda q: math.hypot(q[0] - d[0], q[1] - d[1])); ortho(pen, d, t, 0.75)
        if zone is not None: pen.text(d[0] - 600000, d[1] - 2.6 * r, 1200000, 130000, f"{ds['id']} 门控（开门亮灯）", 6, align=PP_ALIGN.CENTER)
    # board
    if zone is None or V.inside(V.T(AL1)):
        a = V.T(AL1)
        if FLOOR == "FF":
            pen.shape(MSO_SHAPE.OVAL, a[0], a[1], 1.6 * r, 1.6 * r, fill=WHITE, lw=1.25)
            pen.line(a[0] - 0.55 * r, a[1] + 0.55 * r, a[0] + 0.55 * r, a[1] - 0.55 * r, w=1.25)
            pen.text(a[0] - 2400000 - r, a[1] + 1.0 * r, 2400000, 150000, "/".join(WL) + " 由一层 AL1 沿楼梯井引上", 6.5, True, align=PP_ALIGN.RIGHT)
        else:
            board(pen, a[0], a[1], 0.8 * r, "AL1")
    # lamps + switches on top
    for n, pt in LP.items():
        if not inz(n): continue
        x, y = V.T(pt); cn_light(pen, n, x, y, r)
        if n in LAMP: pen.text(x + 1.1 * r, y - 2.3 * r, 2 * r, 1.5 * r, CL[LAMP[n]], 8 if zone else 6.5, True, color=WIRE)
    for sid, pn, loc, sh in PLATES:
        if not swz(pn): continue
        x, y = V.T(MOUNT[pn]); g, keys = PANEL[pn]
        if g == 0:
            pen.box(x, y, "FI", 0.8 * r, 5); pen.text(x + 1.5 * r, y - 2.2 * r, 600000, 150000, sid, 8 if zone else 6.5, True); continue
        two = any(len([pp for kk in next(cc for cc in CIRC if cc[0] == KEY_CIRC[k])[2] for pp in KEY2PANEL.get(kk, [])]) > 1
                  for k in keys if KEY_CIRC.get(k))
        ex, ey = cn_switch(pen, x, y, g, two, 0.8 * r)
        pen.text(ex + 0.2 * r, ey - 1.6 * r, 600000, 150000, sid, 8 if zone else 6.5, True)
    for fid, pn, k in FCUS:
        if FLOOR == "FF" or zone not in (None, 3): continue
        x, y = V.T(MOUNT[pn]); pen.box(x, y, "FCU", 0.8 * r, 5)
    draw_ghosts(pen, V, r, swz, 7 if zone else 6)


# ---------------- 图签 (Chinese title block) ----------------
def frame_cn(pen, title, number, scale="1:100（A3）"):
    pen.shape(MSO_SHAPE.RECTANGLE, W / 2, H / 2, W - 200000, H - 200000, fill=None, lw=1.5)
    pen.text(250000, 180000, 8000000, 300000, title, 15, True)
    tbx, tby, tbw, tbh = RIGHT, 5000000, W - 100000 - RIGHT, 1758000
    pen.shape(MSO_SHAPE.RECTANGLE, tbx + tbw / 2, tby + tbh / 2, tbw, tbh, fill=None, lw=1.0)
    cells = [("工程名称", f"27PR 住宅改造（{FL_CN}）"), ("图名", title), ("图号", number), ("比例", scale),
             ("专业 / 阶段", f"电气 · 施工图（{REV_CN}）"), ("设计 / 校对", "—— / ——（待电工复核签字）")]
    rh = tbh / len(cells)
    for i, (k, v) in enumerate(cells):
        y = tby + i * rh
        if i: pen.line(tbx, y, tbx + tbw, y, w=0.5)
        pen.text(tbx + 50000, y, 700000, rh, k, 7, True, anchor=MSO_ANCHOR.MIDDLE)
        pen.text(tbx + 820000, y, tbw - 870000, rh, v, 7, anchor=MSO_ANCHOR.MIDDLE)
    pen.line(tbx + 780000, tby, tbx + 780000, tby + tbh, w=0.5)


def scale_bar(pen, mpm, n=3):
    for i in range(n):
        pen.shape(MSO_SHAPE.RECTANGLE, 400000 + mpm * (i + 0.5), 6560000, mpm, 40000, fill=INK if i % 2 == 0 else WHITE, lw=0.5)
    pen.text(400000 + mpm * n + 40000, 6500000, 300000, 120000, f"{n} m", 6)


def border(cell, color="808080", w=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag), w=str(w)); sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=color)


def table(slide, x, y, widths, rows, rh, size):
    tbl = slide.shapes.add_table(len(rows), len(widths), Emu(x), Emu(y), Emu(sum(widths)), Emu(rh * len(rows))).table
    tbl.first_row = True; tbl.horz_banding = False
    for j, w in enumerate(widths): tbl.columns[j].width = Emu(w)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Emu(rh)
        for j, v in enumerate(row):
            c = tbl.cell(i, j); c.text = v; c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0xE7, 0xE6, 0xE6) if i == 0 else WHITE
            c.margin_top = c.margin_bottom = Emu(10000); c.margin_left = c.margin_right = Emu(35000)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE; border(c)
            for para in c.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(size); run.font.bold = (i == 0); run.font.name = "Microsoft YaHei"; run.font.color.rgb = INK
    return tbl


# ================= 电施-01 设计说明 · 图例 · 灯具表 =================
s1 = new_slide(); pen = Pen(s1)
frame_cn(pen, f"{FL_CN}照明 设计说明 · 图例 · 灯具表", SN(1), "—")
NOTES = ["设计说明",
         "1. 依据：GB 50034-2013《建筑照明设计标准》、GB/T 50786-2012《建筑电气制图标准》、09DX001《建筑电气工程设计常用图形和文字符号》。",
         "   本工程位于英国，实际施工、器材与验收以 BS 7671 及 Part P 为准；本套图为国标表达方式，供沟通与核对。",
         ("2. 一层照明由照明配电箱 AL1（暂定楼梯下，现场定）引出 3 个回路：WL1 前区、WL2 后区（含卫生间排气扇）、WL3 户外，均带 30 mA 漏电保护。"
          if FLOOR == "GF" else "2. 二层照明由一层 AL1 引出 2 个回路，沿楼梯井引上：WL4 西区、WL5 东区（各含卫生间排气扇），均带 30 mA 漏电保护。"),
         "3. 导线 BV-3×2.5（L、N、PE），穿 PC20 管沿墙 / 顶板暗敷（WC/CC）。线上短斜线为导线根数，未标注者为 3 根。",
         "4. 开关：古铜色 86 型（NAB 系列）暗装，中心距地 1.2 m（国内常用 1.3 m，本工程按英国 1.2 m），距门框 0.15–0.20 m，设在门开启（锁）侧，开门即可开灯。",
         "5. 双控 / 三控：同一字母的灯可在多处开关控制；两处为单极双控开关，三处时中间一处为中途开关（三控），开关间敷设联络线（长虚线）。",
         ("6. 前储物间加门控开关 MK（门框锁侧，开门亮灯、关门灭灯），与 S2 第 2 键并联。" if FLOOR == "GF"
          else "6. 楼梯吊灯 F06 三控：一层 S6（楼梯底）、二层 S13（楼梯顶）为两路开关，二层 S12 第 2 键为中途开关。"),
         "7. 卫生间灯具防护等级 ≥ IP44，户外灯具 ≥ IP65。",
         ("8. FCU1–5 为厨房 / 洗衣房电器带保险丝开关（13A），属插座回路，见插座平面图。" if FLOOR == "GF"
          else "8. 西卫、后卫排气扇由门外 FI 隔离开关控制，风扇随卫生间灯联动（接灯开关负载侧），带延时。"),
         ("9. 按常理补全（Rev B）：l 回路 4 盏筒灯三处控制（S7 为中途开关）；38 号键控洗衣房筒灯；围栏地灯 4 盏；前储物间门控开关。" if FLOOR == "GF"
          else "9. 按常理补全（Rev B）：原图“6+”作为西卫筒灯独立单联开关 S2；主卧主灯位 b 预留（产品待选）。")]
NOTES = CN_NOTES or NOTES
pen.text(250000, 620000, 8100000, 2600000, "\n".join(NOTES), 7.5)
# legend (left-bottom)
r = 55000; y0 = 3500000
pen.text(250000, y0 - 250000, 3000000, 200000, "图例", 9, True)
LEG = [("pend", "吊灯 / 吸顶灯", "按产品", "吊装 / 吸顶"), ("pending", "主灯位（产品待选）", "待选", "吊装"), ("down", "筒灯 GU10", "LED 5W 3000K", "嵌入式"),
       ("wall", "壁灯 / 户外墙灯", "按产品", "壁装 1.8–2.0 m"), ("ground", "庭院地埋 / 围栏灯", "LED IP65", "地面"),
       ("sw", "单极开关（短线数 = 联数）", "86 型 10A", "暗装 1.2 m"), ("sw2", "单极双控开关", "86 型 10A", "暗装 1.2 m"),
       ("mk", "门控开关 MK", "门框式", "门框锁侧"), ("fi", "排气扇隔离开关 FI", "86 型", "暗装"), ("al", "照明配电箱 AL1", "见系统图", "暗装 1.5 m（底边）")]
if FLOOR == "FF": LEG = [l for l in LEG if l[0] not in ("mk", "wall", "ground")]
if not HAS_FI: LEG = [l for l in LEG if l[0] != "fi"]
for i, (k, a, b, c) in enumerate(LEG):
    col, row = i % 2, i // 2; x = 400000 + col * 4100000; y = y0 + row * 300000
    if k == "pend": pen.pendant(x, y, r)
    elif k == "pending": pen.pendant(x, y, r, True)
    elif k == "down": pen.down(x, y, r)
    elif k == "wall": cn_light(pen, "壁灯 x", x, y, r)
    elif k == "ground": pen.ground(x, y, r)
    elif k == "sw": cn_switch(pen, x - 40000, y + 40000, 2, False, r)
    elif k == "sw2": cn_switch(pen, x, y + 20000, 1, True, r)
    elif k == "mk": pen.box(x, y, "MK", r, 5)
    elif k == "fi": pen.box(x, y, "FI", r, 5)
    elif k == "al": board(pen, x, y, 0.8 * r, "")
    pen.text(x + 250000, y - 80000, 3700000, 170000, f"{a}　｜ {b}　｜ {c}", 7)
# lamp schedule (right)
cnt = {}
for n in LP:
    if n in LAMP:
        key = (lamp_type(n), n.split()[1].split("_")[0] if lamp_type(n) != "筒灯" else "GU10")
        cnt[key] = cnt.get(key, 0) + 1
rows = [["编号", "名称", "光源 / 功率", "防护", "数量"]]
PROD = {"G06": "G06 繁花吊灯 Φ80", "G07": "G07 繁花吊灯 Φ60", "G08": "G08 单头小吊灯", "G15": "G15 法式吊灯", "G16": "G16 吸顶灯",
        "F32": "F32 黑胡桃壁灯", "F06": "F06 贝壳吊灯", "F11": "F11 玄关吸顶灯", "F19": "F19 吸顶灯", "F20": "F20 吸顶灯",
        "F22": "F22 云朵吸顶灯", "F23": "F23 海派吸顶灯", "G25": "G25 户外入户灯", "G26": "G26 花园壁灯", "GroundLED": "围栏地灯", "GU10": "GU10 筒灯", "主灯待选": "主灯（待选）"}
for i, ((t, code), q) in enumerate(sorted(cnt.items(), key=lambda kv: kv[0][1]), 1):
    ip = "IP65" if code in ("G25", "G26", "GroundLED") else ("IP44（卫生间）/IP20" if code == "GU10" else "IP20")
    src_ = "LED 5W 3000K" if code == "GU10" else ("LED 3000K" if code == "GroundLED" else "按产品")
    rows.append([f"D{i}", PROD.get(code, code), src_, ip, str(q)])
pen.text(RIGHT + 60000, 300000, 3300000, 200000, "灯具表", 9, True)
table(s1, RIGHT + 60000, 560000, (380000, 1300000, 850000, 650000, 300000), rows, 300000, 6.5)

# ================= 电施-02 一层照明平面图（全层） =================
s2 = new_slide(); pen = Pen(s2)
V = View((40, 120, 1990, 1110), (300000, 850000, RIGHT - 200000, 6650000))
mpm = V.plan(s2)
frame_cn(pen, f"{FL_CN}照明平面图", SN(2))
draw_plan(s2, pen, V, None, 0.09 * mpm)
pen.text(RIGHT + 60000, 300000, 3300000, 2000000,
         f"说明：\n1. 灯旁字母 = 控制它的开关键，见{SN(4 + NZ)}。\n2. 粗实线 = AL1 引出回路 {'/'.join(WL)}（先到开关盒）；细实线 = 同一字母灯具之间；点线 = 开关到所控灯；长虚线 = 双控联络线。\n"
         f"3. 线上短斜线 = 导线根数，未标注为 3 根。\n4. 分区放大见{SN(3)}~{SN(2 + NZ)}。\n5. S# = 开关编号；MK = 门控开关" + ("；FI = 排气扇隔离开关。" if HAS_FI else "。"), 7)
for z, d in ZONES.items():
    a, b, c, e = d["view"]; (x0, y0), (x1, y1) = V.T((a * U, b * U)), V.T((c * U, e * U))
    pen.shape(MSO_SHAPE.RECTANGLE, (x0 + x1) / 2, (y0 + y1) / 2, x1 - x0, y1 - y0, fill=None, lw=1.0, color=RED, dash=MSO_LINE_DASH_STYLE.LONG_DASH)
    pen.text(x0 + 40000, y0 + 30000, 2000000, 150000, SN(z + 2), 7, True, color=RED)
scale_bar(pen, mpm)

# ================= 电施-03~05 分区放大 =================
for z, d in ZONES.items():
    sl = new_slide(); pen = Pen(sl)
    Vz = View(d["view"], (300000, 850000, RIGHT - 200000, 6650000))
    mz = Vz.plan(sl)
    bx0, by0, bx1, by1 = Vz.box
    for (mx0, my0, mx1, my1) in ((0, 0, W, by0), (0, by1, W, H), (0, by0, bx0, by1), (bx1, by0, W, by1)):
        pen.shape(MSO_SHAPE.RECTANGLE, (mx0 + mx1) / 2, (my0 + my1) / 2, mx1 - mx0, my1 - my0, fill=WHITE, lw=0)
    frame_cn(pen, f"{FL_CN}照明平面放大图（{d['name']}）", SN(z + 2), "1:50（A3）")
    draw_plan(sl, pen, Vz, z)
    y = 300000
    pen.text(RIGHT + 60000, y, 3300000, 180000, "本图开关（键 1→键 4 从左到右）", 8.5, True); y += 230000
    for sid, pn, loc, sh in PLATES:
        if sh != z: continue
        g, keys = PANEL[pn]
        ks = "  ".join(f"键{i + 1}={CL[KEY_CIRC[k]] if KEY_CIRC.get(k) else EXTRA_TAG.get(k, '?')}" for i, k in enumerate(keys))
        pen.text(RIGHT + 60000, y, 3300000, 150000, f"{sid}  {loc}", 7, True); y += 160000
        pen.text(RIGHT + 260000, y, 3100000, 150000, ks, 7, color=WIRE); y += 190000
    for ds in DOORSW:
        if ds["sheet"] != z: continue
        pen.text(RIGHT + 60000, y, 3300000, 150000, f"{ds['id']}  {ds['label']} → {CL[ds['circuit']]}（与 {ds['parallel']} 并联）", 7, True); y += 200000
    scale_bar(pen, mz, 2)

# ================= 电施-06 AL1 系统图 =================
s6 = new_slide(); pen = Pen(s6)
frame_cn(pen, "AL1 照明配电箱系统图（一、二层）", SN(3 + NZ), "—")
WL_ALL = dict(WL_GF_DEF); WL_ALL.update(WL_FF)
step = min(700000, 3700000 // (len(WL_ALL) + 1))   # breaker rows + 1 spare must fit above the notes
bx = 2000000; top = 1300000; bot = top + 450000 + len(WL_ALL) * step
pen.line(bx, top, bx, bot, w=3.0)
pen.text(400000, 1450000, 3000000, 300000, "进线：引自总配电箱 / 电表箱\n（英国：Consumer Unit 内主开关 + RCBO）", 8)
pen.line(1000000, 1300000, bx, 1300000, w=1.5)
pen.text(bx + 100000, top - 260000, 1500000, 200000, "AL1  暗装，底边距地 1.5 m", 8, True)
for i, (w, d) in enumerate(WL_ALL.items()):
    y = top + 450000 + i * step
    pen.line(bx, y, bx + 900000, y, w=1.0)
    pen.shape(MSO_SHAPE.RECTANGLE, bx + 1100000, y, 400000, 220000, fill=WHITE, lw=1.0)
    pen.line(bx + 1300000, y, bx + 4200000, y, w=1.0)
    n_l = d.get("n") or sum(len(next(cc for cc in CIRC if cc[0] == c)[3]) for c in d["circ"])
    pen.text(bx + 900000, y - 330000, 1600000, 150000, d["br"], 7, True)
    pen.text(bx + 1450000, y - 190000, 2700000, 150000, "BV-3×2.5  PC20  WC/CC", 7)
    fl = "二层" if w in WL_FF else "一层"
    lets = ' '.join(CL[c] for c in d['circ']) if w in WL else "见" + fl + "图"
    pen.text(bx + 4300000, y - 150000, 3500000, 330000, f"{w}  {d['name']}\n灯 {n_l} 盏（{fl}，字母 {lets}）", 7)
y = bot
pen.line(bx, y, bx + 900000, y, w=1.0, dash=MSO_LINE_DASH_STYLE.DASH)
pen.text(bx + 1000000, y - 150000, 5000000, 300000, "备用 1 路（C16 RCBO）", 7.5)
pen.text(400000, 5550000, 7900000, 1100000,
         "说明：\n1. 照明回路按区域划分，每路一个 30 mA RCBO（见上表回路名称）；单回路灯具数量与功率均在允许范围内（灯具均为 LED）。插座与专线回路见网页「配电箱」页。\n"
         "2. 英国规范（BS 7671）照明回路常用 1.5 mm² 双芯加地线 + B6 RCBO；国内常用 BV-2.5 + C16。施工由电工按当地规范定。\n" +
         ("3. 排气扇（FI）接所在区的照明回路；" if HAS_FI else "3. 卫生间排气扇无单独开关，接所在卫生间灯回路；") + f"楼梯吊灯 F06 属 {next((w for w, d in WL_FF.items() if '楼梯' in d['name']), '二层回路')}，一层 S6 只做三控开关点。\n4. " + "、".join(WL_FF) + " 沿楼梯井暗管引上至二层。", 7.5)

# ================= 电施-07 开关面板表 + 回路表 =================
s7 = new_slide(); pen = Pen(s7)
frame_cn(pen, f"{FL_CN}开关面板表 · 照明回路表", SN(4 + NZ), "—")
rows = [["编号", "位置", "规格", "键 1", "键 2", "键 3", "键 4"]]
for sid, pn, loc, sh in PLATES:
    g, keys = PANEL[pn]
    cells = []
    for k in keys:
        c = KEY_CIRC.get(k)
        if c:
            n_ctrl = len([pp for kk in next(cc for cc in CIRC if cc[0] == c)[2] for pp in KEY2PANEL.get(kk, [])]) + len(REMOTE.get(c, []))
            way = {1: "单控", 2: "双控", 3: "三控"}.get(n_ctrl, f"{n_ctrl}控")
            if (sid, k) in MIDDLE_KEYS: way = "中途"
            cells.append(f"{CL[c]} {DESC[c]}（{way}）")
        else:
            cells.append(EXTRA_KEYS.get(k, ""))
    cells += [""] * (4 - len(cells))
    spec = {0: "FI 隔离开关", 1: "一位", 2: "二位", 3: "三位", 4: "四位"}[g]
    rows.append([sid, loc, spec, *cells])
for ds in DOORSW: rows.append([ds["id"], ds["location"], "门控开关", ds["schedule"], "", "", ""])
table(s7, 250000, 650000, (380000, 1250000, 650000, 1550000, 1550000, 1550000, 1450000), rows, 205000, 6)
rows2 = [["字母", "控制灯具", "开关", "回路"]]
for cid, desc, keys, lights in CIRC:
    sw = "，".join(f"{PID[pn]}-{PANEL[pn][1].index(k) + 1}" for k in keys for pn in KEY2PANEL.get(k, []))
    sw += "".join("，" + ds["id"] for ds in DOORSW if ds["circuit"] == cid) + "".join("，" + x for x in REMOTE.get(cid, []))
    rows2.append([CL[cid], DESC[cid], sw, C2W[cid]])
table(s7, RIGHT + 60000, 300000, (300000, 1650000, 900000, 400000), rows2, 205000, 6)

# keep only the new sheets
ids = list(p.slides._sldIdLst)
for sid_el in ids[:N0]:
    p.part.drop_rel(sid_el.rId); p.slides._sldIdLst.remove(sid_el)
p.save(DST)
print("saved", DST, len(p.slides))
