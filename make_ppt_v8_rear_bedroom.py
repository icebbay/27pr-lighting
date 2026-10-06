"""PPT v8 = v6 照明控制图_专业版 + 二层后房间改卧室 4 + 东侧套卫（2026-10-06）.

v6 is the source of the lighting / power tool (FF control plan slide 6, FF socket page slide 7). Its FF wall groups use the
same child frame as v4 slide 6, so the wall edit of ppt_v7_rear_bedroom_walls.py applies unchanged. The user's moved points
(v7 slides 6 / 7 / 8) are copied onto v6 slides 5 / 7 / 8 by shape id (ids are identical in v4 and v6).

    python products/make_ppt_v8_rear_bedroom.py
Writes Downloads/…_v8_照明控制图_专业版_卧室套卫.pptx and a copy in products/source_pptx/.
"""
import copy, os, shutil, sys
from pptx import Presentation
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from ppt_ff_transform import b2s, s2b

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
ns = {'a': A}
DL = os.path.join(os.path.expanduser("~"), "Downloads")
SRC = os.path.join(HERE, "source_pptx", "Blender模型户型图_灯_插座_水路_可编辑_v6_照明控制图_专业版.pptx")
USER = os.path.join(DL, "Blender模型户型图_灯_插座_水路_可编辑_v7_卧室套卫.pptx")
NAME = "Blender模型户型图_灯_插座_水路_可编辑_v8_照明控制图_专业版_卧室套卫.pptx"
DST = os.path.join(DL, NAME)

PX0, PX1 = 3.27, 3.38
SY0, SY1 = 8.96, 9.07
D0, D1 = 9.12, 9.78
WALLS = {
    'V39_Partition_Bed_Bath_W': [(PX0, D1), (PX1, D1), (PX1, 11.51), (PX0, 11.51)],
    'V39_Partition_Bed_Bath_S': [(PX0, SY0), (4.29, SY0), (4.29, SY1), (PX1, SY1), (PX1, D0), (PX0, D0)],
}
LEAF = {'V39_DoorLeaf_Bath': [(PX0 - (D1 - D0), D1 - 0.04), (PX0, D1 - 0.04), (PX0, D1), (PX0 - (D1 - D0), D1)]}
ROOM = "卧室 4"


def set_poly(el, pts, name, nid):
    S = [b2s(x, y) for x, y in pts]
    xs = [p[0] for p in S]; ys = [p[1] for p in S]
    ox, oy = int(min(xs)), int(min(ys))
    w, h = max(1, int(max(xs)) - ox), max(1, int(max(ys)) - oy)
    c = el.find('.//{*}cNvPr'); c.set('name', '墙_' + name); c.set('id', str(nid))
    x = el.find('.//a:xfrm', ns)
    x.find('a:off', ns).set('x', str(ox)); x.find('a:off', ns).set('y', str(oy))
    x.find('a:ext', ns).set('cx', str(w)); x.find('a:ext', ns).set('cy', str(h))
    pl = el.find('.//a:pathLst', ns)
    for ch in list(pl): pl.remove(ch)
    path = etree.SubElement(pl, f'{{{A}}}path', w=str(w), h=str(h))
    for i, (sx, sy) in enumerate(S):
        m = etree.SubElement(path, f'{{{A}}}' + ('moveTo' if i == 0 else 'lnTo'))
        etree.SubElement(m, f'{{{A}}}pt', x=str(int(sx) - ox), y=str(int(sy) - oy))
    etree.SubElement(path, f'{{{A}}}close')


def center_to(sh, x, y):
    cx, cy = s2b(sh.left + sh.width / 2, sh.top + sh.height / 2)
    sx0, sy0 = b2s(cx, cy); sx1, sy1 = b2s(x, y)
    sh.left = int(sh.left + sx1 - sx0); sh.top = int(sh.top + sy1 - sy0)


def groups(shapes, prefix):
    for sh in shapes:
        if sh.shape_type == 6:
            if sh.name.startswith(prefix): yield sh
            else: yield from groups(sh.shapes, prefix)


p = Presentation(SRC)
nid = 100000
for si in (4, 5, 6, 7):                    # v6 slides 5-8 = FF
    s = p.slides[si]
    for g in groups(s.shapes, '墙体'):
        by = {sh.name: sh for sh in g.shapes}
        part = by['墙_V35_Partition_Study_Bath']._element; leaf = by['墙_V35_DoorLeaf_Bath245']._element
        for n, pts in WALLS.items():
            e = copy.deepcopy(part); set_poly(e, pts, n, nid); nid += 1; part.addnext(e)
        for n, pts in LEAF.items():
            e = copy.deepcopy(leaf); set_poly(e, pts, n, nid); nid += 1; leaf.addnext(e)
        g._element.remove(part); g._element.remove(leaf)
    for g in groups(s.shapes, '房间名'):
        for sh in g.shapes:
            if not sh.has_text_frame: continue
            t = sh.text_frame.text
            cx, cy = s2b(sh.left + sh.width / 2, sh.top + sh.height / 2)
            if t == '书房' and cx > 2.0 and cy > 8:
                sh.text_frame.paragraphs[0].runs[0].text = ROOM; center_to(sh, 2.25, 9.2)
            elif t == '卫生间' and abs(cx - 2.73) < 0.1 and abs(cy - 11.03) < 0.1:
                sh.text_frame.word_wrap = True
                h0 = sh.height; sh.height = int(h0 * 3.2); sh.width = int(h0 * 0.9); center_to(sh, 3.53, 10.25)

# user's moved points (v7 slide -> v6 slide), plus the two lights on the control plan (slide 6, inside 二层平面（横向）)
u = Presentation(USER)
MOVED = {6: (5, [243, 252, 378, 384]), 7: (7, [363, 365, 367, 369, 371])}
WATER = [274, 277, 281, 284, 287, 290, 293, 296, 299]
MOVED[8] = (8, WATER)
for us, (vs, ids) in MOVED.items():
    U = {sh.shape_id: sh for sh in u.slides[us - 1].shapes}
    V = {sh.shape_id: sh for sh in p.slides[vs - 1].shapes}
    for i in ids:
        assert V[i].name == U[i].name, (vs, i, V[i].name, U[i].name)
        V[i].left, V[i].top = U[i].left, U[i].top
U6 = {sh.shape_id: sh for sh in u.slides[5].shapes}
plan = [x for x in p.slides[5].shapes if x.name == '二层平面（横向）'][0]
for sh in plan.shapes:
    if sh.shape_id in (243, 252):
        sh.left, sh.top = U6[sh.shape_id].left, U6[sh.shape_id].top

p.save(DST)
shutil.copy(DST, os.path.join(HERE, "source_pptx", NAME))
print("saved", DST)
